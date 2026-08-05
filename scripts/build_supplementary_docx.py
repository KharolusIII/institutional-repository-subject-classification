"""Build a reviewer-friendly DOCX from the public aggregate artifacts."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path("paper_artifacts")
TABLES = ROOT / "tables"
V3 = ROOT / "runs" / "v3_main"
AUDIT = ROOT / "runs" / "convergence_audit"


def read_rows(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def shade(cell, fill: str):
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def add_table(document: Document, headers: list[str], values: list[list[str]], widths=None):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Shading Accent 1"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header
        shade(cell, "D9EAF7")
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for value_row in values:
        cells = table.add_row().cells
        for index, value in enumerate(value_row):
            cells[index].text = str(value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_picture(document: Document, path: Path, caption: str, width=6.7):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    caption_paragraph = document.add_paragraph(caption)
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.style = document.styles["Caption"]


def metric(row, key, digits=4):
    return f"{float(row[key]):.{digits}f}"


def configuration_name(row):
    return " / ".join(
        str(row.get(key, ""))
        for key in ("representation", "feature_set", "preprocessing", "classifier")
        if row.get(key, "")
    )


def model_display_name(value):
    return {
        "sbert_frozen": "Fixed DistilUSE embeddings",
        "labse_frozen": "Fixed LaBSE embeddings",
        "sbert_finetuned": "DistilUSE-backbone sequence classifier",
        "labse_finetuned": "LaBSE-backbone sequence classifier",
    }.get(str(value), str(value))


def build(output: Path):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(20)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Supplementary Material\n")
    subtitle = title.add_run("Beyond Embeddings: Language-Aware Sparse Representations and\n"
                             "Fine-Tuned Multilingual Transformers")
    subtitle.font.size = Pt(13)
    subtitle.italic = True
    note = document.add_paragraph(
        "Public reviewer package · aggregate evidence only · generated from the versioned artifacts"
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_heading("S1. Evidence runs and public-data boundary", level=1)
    document.add_paragraph(
        "This supplement documents two frozen runs. It contains aggregate statistics and figures only. "
        "Repository texts, handles, item-level predictions, split manifests, model weights, caches, "
        "credentials, and personal paths are excluded."
    )
    add_table(document, ["Run", "Purpose", "Source commit"], [
        ["20260802T165746Z_full_20k_v3_segmented_finetuned_2026",
         "Main 233-candidate experiment", "7181fcceb7a6"],
        ["20260804T102544Z_convergence_audit_20k_2026",
         "Eight-fit LinearSVC convergence audit", "6e46e5c2af02"],
    ], [2.6, 2.3, 1.2])

    document.add_heading("S2. Corpus, coverage, and frozen split", level=1)
    dataset = read_rows(V3 / "dataset_statistics.csv")[0]
    units = read_rows(V3 / "text_unit_statistics.csv")[0]
    coverage = read_rows(V3 / "fulltext_coverage.csv")[0]
    add_table(document, ["Property", "Value"], [
        ["Documents", dataset["N_documents"]],
        ["Labels / assignments", f'{dataset["N_labels"]} / {dataset["N_assignments"]}'],
        ["Abstract segments / full-text files", f'{units["N_abstract_segments"]} / {units["N_fulltext_files"]}'],
        ["Items with multiple abstracts / full texts", f'{units["items_with_multiple_abstracts"]} / {units["items_with_multiple_fulltexts"]}'],
        ["Represented full-text characters", f'{float(coverage["represented_character_percentage"]):.2f}%'],
        ["Items reaching 400k cap", f'{coverage["truncated_documents"]} ({float(coverage["truncated_document_percentage"]):.2f}%)'],
        ["Train / calibration / validation / test", "12,836 / 1,927 / 1,988 / 2,959"],
    ], [4.5, 2.0])
    add_picture(document, ROOT / "figures" / "dataset_construction_workflow.png",
                "Figure S1. Dataset construction and leakage-controlled partitioning.")

    document.add_heading("S3. Language audit", level=1)
    agreement = read_rows(V3 / "language_agreement_summary.csv")
    gt = read_rows(V3 / "abstract_language_ground_truth_summary.csv")[0]
    add_table(document, ["Evaluation", "N", "Agreement / accuracy"], [
        ["Declared vs. detected abstract", agreement[0]["auditable_N"], f'{100*float(agreement[0]["agreement_rate"]):.2f}%'],
        ["Declared vs. detected full text", agreement[1]["auditable_N"], f'{100*float(agreement[1]["agreement_rate"]):.2f}%'],
        ["Historical audited abstract reference", gt["matched_segments"], f'{100*float(gt["accuracy"]):.2f}%'],
    ], [3.5, 1.2, 1.7])
    document.add_paragraph(
        "The historical reference contains 23,055 human/historical-detector agreements and 554 "
        "manually audited disagreements; it is not an independently manual annotation set."
    )

    document.add_heading("S4. Experimental protocol", level=1)
    document.add_paragraph(
        "The primary validation set contains 233 candidates: 189 sparse configurations, 42 fixed "
        "sentence-embedding configurations, and two supervised Transformer sequence classifiers. "
        "Macro-F1 is the prespecified primary "
        "selection criterion; thresholds are estimated on the dedicated calibration partition."
    )
    add_picture(document, ROOT / "figures" / "experimental_protocol_workflow.png",
                "Figure S2. Calibrated multilingual long-document evaluation protocol.")
    document.add_paragraph(
        "Dense-model terminology: historical sbert* identifiers denote DistilUSE. Fixed DistilUSE "
        "and LaBSE candidates use their complete, unchanged Sentence-Transformers pipelines. The "
        "supervised candidates instead initialize AutoModelForSequenceClassification from the "
        "Transformer module distributed with each checkpoint, create a new 37-output head, and "
        "jointly optimize backbone and head with class-weighted BCE. This is supervised "
        "sequence-classification fine-tuning, not contrastive sentence-embedding fine-tuning."
    )

    document.add_heading("S5. Validation and family-level test comparison", level=1)
    leaders = read_rows(TABLES / "validation_leaders.csv")
    add_table(document, ["Selection basis", "Configuration", "Macro-F1", "Micro-F1"], [
        [row["selection_basis"],
         f'{row["representation"]} / {row["feature_set"]} / {row["classifier"]}',
         f'{float(row["f1_macro"]):.4f}', f'{float(row["f1_micro"]):.4f}']
        for row in leaders
    ], [1.1, 3.4, 1.0, 1.0])

    validation_grid = read_rows(V3 / "results_validation_grid.csv")
    validation_grid.sort(key=lambda row: float(row["f1_macro"]), reverse=True)
    add_table(document, ["Rank", "Validation configuration", "Macro-F1", "Micro-F1"], [
        [str(index), configuration_name(row), metric(row, "f1_macro"), metric(row, "f1_micro")]
        for index, row in enumerate(validation_grid[:15], start=1)
    ], [0.55, 4.35, 0.9, 0.9])
    document.add_paragraph(
        "Table S6 reports the 15 highest Macro-F1 configurations. The complete 233-row validation "
        "grid is available at https://github.com/KharolusIII/"
        "institutional-repository-subject-classification/blob/main/paper_artifacts/runs/v3_main/"
        "results_validation_grid.csv."
    )

    family = read_rows(TABLES / "family_test_comparison.csv")
    add_table(document, ["Family", "Selected source", "Val. Macro-F1", "Test Macro-F1", "Test Micro-F1"], [
        [model_display_name(row["evaluation_family"]), row["feature_set"],
         f'{float(row["validation_f1_macro"]):.4f}', f'{float(row["f1_macro"]):.4f}',
         f'{float(row["f1_micro"]):.4f}'] for row in family
    ], [1.25, 1.65, 1.2, 1.2, 1.2])

    document.add_heading("S6. Convergence audit and final held-out metrics", level=1)
    convergence = read_rows(AUDIT / "results_validation.csv")
    convergence.sort(key=lambda row: float(row["f1_macro"]), reverse=True)
    add_table(document, ["Configuration", "Macro-F1", "Micro-F1"], [
        [configuration_name(row), metric(row, "f1_macro"), metric(row, "f1_micro")]
        for row in convergence
    ], [4.7, 1.0, 1.0])
    final = read_rows(TABLES / "final_test_metrics.csv")[0]
    intervals = {row["metric"]: row for row in read_rows(AUDIT / "bootstrap_ci.csv")}
    add_table(document, ["Metric", "Estimate", "95% CI where applicable"], [
        ["Subset accuracy", metric(final, "subset_accuracy"), "—"],
        ["Hamming loss", metric(final, "hamming_loss"), "—"],
        ["Macro-F1", metric(final, "f1_macro"),
         f'[{float(intervals["f1_macro"]["ci_lower"]):.4f}, {float(intervals["f1_macro"]["ci_upper"]):.4f}]'],
        ["Micro-F1", metric(final, "f1_micro"),
         f'[{float(intervals["f1_micro"]["ci_lower"]):.4f}, {float(intervals["f1_micro"]["ci_upper"]):.4f}]'],
        ["Recall@3", metric(final, "recall_at_3"), "—"],
        ["Precision@3", metric(final, "precision_at_3"), "—"],
    ], [2.3, 1.4, 2.8])
    document.add_paragraph(
        "Both C=0.5 tolerance settings converged for all 37 labels and tied on all stored validation "
        "metrics. The stored tol=0.0005 setting is not claimed to outperform tol=0.0001."
    )
    add_picture(document, ROOT / "figures" / "test_metrics_overview.png",
                "Figure S3. Final held-out metric overview.", width=6.2)

    paired = read_rows(TABLES / "paired_bootstrap_updated.csv")
    add_table(document, ["Comparison", "Metric", "Difference", "95% CI", "P(candidate > BM25)"], [
        [f'{model_display_name(row["evaluation_family"])} minus final BM25', row["metric"],
         f'{float(row["difference_candidate_minus_final_bm25"]):.4f}',
         f'[{float(row["ci_lower"]):.4f}, {float(row["ci_upper"]):.4f}]',
         f'{float(row["probability_candidate_better"]):.4f}']
        for row in paired
    ], [2.3, 0.9, 1.1, 1.6, 0.7])

    document.add_heading("S7. Supervised sequence-classification trajectories", level=1)
    transformer = read_rows(TABLES / "transformer_validation.csv")
    add_table(document, ["Model", "Selected epoch", "Val. Macro-F1", "Val. Micro-F1"], [
        [model_display_name(row["model"]), row["selected_epoch"], metric(row, "f1_macro"), metric(row, "f1_micro")]
        for row in transformer
    ], [2.2, 1.3, 1.4, 1.4])

    document.add_heading("S8. Complete per-subject final results", level=1)
    per_label = read_rows(TABLES / "per_label_test.csv")
    per_label.sort(key=lambda row: float(row["f1"]), reverse=True)
    add_table(document, ["Subject", "Test support", "Precision", "Recall", "F1", "AP"], [
        [row["label"], row["support_test"], metric(row, "precision"), metric(row, "recall"),
         metric(row, "f1"), metric(row, "average_precision")] for row in per_label
    ], [2.5, 0.8, 0.8, 0.8, 0.8, 0.8])
    add_picture(document, ROOT / "figures" / "per_label_f1.png",
                "Figure S4. Final held-out F1 by subject.", width=6.7)

    document.add_heading("S9. Diagnostic figures", level=1)
    add_picture(document, ROOT / "figures" / "per_label_confusion_matrices.png",
                "Figure S5. One-vs-rest confusion matrices for all 37 subjects.", width=6.7)
    add_picture(document, ROOT / "figures" / "true_vs_predicted_support.png",
                "Figure S6. True versus predicted support by subject.", width=6.7)
    add_picture(document, ROOT / "figures" / "label_error_counts.png",
                "Figure S7. False-positive and false-negative counts by subject.", width=6.7)
    add_picture(document, ROOT / "figures" / "support_vs_f1.png",
                "Figure S8. Relationship between test support and subject-level F1.", width=6.2)
    add_picture(document, ROOT / "figures" / "label_cooccurrence_heatmap.png",
                "Figure S9. Label co-occurrence structure in the evaluation cohort.", width=6.7)

    document.add_heading("S10. Artifact index and reproduction", level=1)
    document.add_paragraph(
        "The GitHub paper_artifacts directory includes the complete 233-row validation export, "
        "eight-row convergence audit, per-label convergence diagnostics, calibrated thresholds, "
        "paired-bootstrap differences, language distributions, label coverage, confusion matrices, "
        "error summaries, learning curves, and publication figures. The source configurations are "
        "configs/run_full_20k_v3_2026.yaml and configs/run_convergence_audit_20k_2026.yaml."
    )
    document.add_paragraph(
        "Code and aggregate evidence: https://github.com/KharolusIII/"
        "institutional-repository-subject-classification"
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(output)


if __name__ == "__main__":
    build(ROOT / "CACIC_2026_Supplementary_Material.docx")
